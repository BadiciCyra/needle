"""Program yöneticisi raporu: akış, sektörler, havuzdaki eksik yetkinlikler, uygun bulunamayan ihtiyaçlar.

Rapor her istekte veritabanından hesaplanır (ayrı bir tablo yok); Excel çıktısı aynı veriden üretilir.

Eksik yetkinlik listesi iki kaynaktan beslenir:
- "yakındı ama" adaylarının eksik yetkinliği (LLM gerekçesindeki missing_capability)
- havuzda doğrudan çözen girişim bulunamayan ihtiyaçların aranan yetkinlikleri
Aynı ihtiyaç bir ifadeyi bir kez sayar; anlamca aynı ifadeler embedding benzerliğiyle tek satırda toplanır.
"""

from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import embedder_dep
from app.auth import require_admin
from app.config import Settings, get_settings
from app.db.models import (
    Application,
    BriefRecord,
    Introduction,
    Match,
    MatchRun,
    OpenCall,
    Pilot,
    Startup,
    User,
)
from app.db.session import get_db
from app.embeddings import Embedder, cosine

router = APIRouter(prefix="/admin", tags=["program yöneticisi raporu"])

SIMILAR_PHRASE = 0.8


class Funnel(BaseModel):
    needs: int = Field(description="Girilen ihtiyaç")
    briefed: int = Field(description="Brief'i tamamlanan")
    matched: int = Field(description="Eşleştirmesi yapılan")
    no_match: int = Field(description="Son eşleştirmede doğrudan çözen girişim bulunamayan")
    introduced: int = Field(description="En az bir tanıştırma isteği giden")
    piloted: int = Field(description="En az bir pilotu açılan")
    worked: int = Field(description="Pilot sonucu 'evet' ya da 'kısmen' olan")


class IntroStats(BaseModel):
    total: int
    waiting: int
    accepted: int
    declined: int
    acceptance_rate: float | None = Field(description="Cevaplananlar içinde kabul oranı")
    avg_response_days: float | None
    via_admin: int = Field(description="Hesabı olmayan girişim adına yöneticinin cevapladığı")


class PilotStats(BaseModel):
    total: int
    active: int
    stale: int
    done: int
    result_evet: int
    result_kismen: int
    result_hayir: int


class SectorRow(BaseModel):
    sector: str
    needs: int
    no_match: int
    introduced: int
    pilots: int
    worked: int


class MissingCapability(BaseModel):
    capability: str
    needs: int = Field(description="Bu yetkinliği arayıp havuzda bulamayan ihtiyaç sayısı")
    variants: list[str] = Field(description="Aynı satırda toplanan farklı yazımlar")
    need_titles: list[str]


class UnmetNeed(BaseModel):
    brief_id: int
    title: str
    organization: str | None
    sector: str
    required_capabilities: list[str]
    open_call_id: int | None
    applications: int


class PoolRow(BaseModel):
    sector: str
    startups: int
    with_account: int = Field(description="Doğrulanmış hesabı olan")


class ReportOut(BaseModel):
    generated_at: datetime
    funnel: Funnel
    introductions: IntroStats
    pilots: PilotStats
    sectors: list[SectorRow]
    missing_capabilities: list[MissingCapability]
    unmet_needs: list[UnmetNeed]
    pool: list[PoolRow]
    open_calls: int
    applications: int


def _tr_lower(text: str) -> str:
    return " ".join(text.replace("İ", "i").replace("I", "ı").lower().split())


def _sector_label(text: str | None) -> str:
    """Firma profili "Enerji", LLM "enerji" yazabilir: ilk harf Türkçe kurala göre büyütülür."""
    text = (text or "").strip()
    if not text:
        return "Belirtilmemiş"
    return {"i": "İ", "ı": "I"}.get(text[0], text[0].upper()) + text[1:]


def group_phrases(
    items: list[tuple[str, str]], embedder: Embedder, threshold: float = SIMILAR_PHRASE
) -> list[MissingCapability]:
    """(ifade, ihtiyaç başlığı) çiftlerini anlamca gruplar. Saf fonksiyon: veritabanı gerektirmez.

    Önce yazımı aynı olanlar (Türkçe küçük harf) birleşir, sonra en sık geçen yazımdan başlayarak
    embedding benzerliği eşiği geçenler aynı gruba katılır. Bir ihtiyaç bir grubu en fazla bir kez sayar.
    """
    by_form: dict[str, dict] = {}
    for phrase, need in items:
        key = _tr_lower(phrase)
        if not key:
            continue
        entry = by_form.setdefault(key, {"spellings": Counter(), "needs": set()})
        entry["spellings"][phrase.strip()] += 1
        entry["needs"].add(need)
    if not by_form:
        return []

    forms = sorted(by_form, key=lambda k: (-len(by_form[k]["needs"]), k))
    vectors = dict(zip(forms, embedder.embed_documents(forms)))
    groups: list[dict] = []
    for form in forms:
        target = next((g for g in groups if cosine(vectors[g["forms"][0]], vectors[form]) >= threshold), None)
        if target is None:
            groups.append({"forms": [form], "needs": set(by_form[form]["needs"])})
        else:
            target["forms"].append(form)
            target["needs"] |= by_form[form]["needs"]

    out = []
    for g in groups:
        spellings = Counter()
        for form in g["forms"]:
            spellings.update(by_form[form]["spellings"])
        ordered = [s for s, _ in spellings.most_common()]
        out.append(
            MissingCapability(
                capability=ordered[0], needs=len(g["needs"]), variants=ordered[1:], need_titles=sorted(g["needs"])
            )
        )
    return sorted(out, key=lambda m: (-m.needs, m.capability))


def build_report(session: Session, embedder: Embedder, settings: Settings) -> ReportOut:
    now = datetime.now(timezone.utc)
    briefs = session.scalars(select(BriefRecord).order_by(BriefRecord.id)).all()

    latest_run = {
        brief_id: run_id
        for brief_id, run_id in session.execute(select(MatchRun.brief_id, func.max(MatchRun.id)).group_by(MatchRun.brief_id))
    }
    shortlisted_runs = set(
        session.scalars(select(Match.run_id).where(Match.kind == "shortlist", Match.run_id.in_(latest_run.values())))
    )
    intros = session.scalars(select(Introduction)).all()
    pilots = session.scalars(select(Pilot)).all()
    calls = {c.brief_id: c for c in session.scalars(select(OpenCall))}
    app_counts = dict(session.execute(select(Application.call_id, func.count()).group_by(Application.call_id)).all())

    intro_briefs = {i.brief_id for i in intros}
    pilots_by_brief: dict[int, list[Pilot]] = defaultdict(list)
    for p in pilots:
        pilots_by_brief[p.brief_id].append(p)

    def sector_of(record: BriefRecord) -> str:
        organization = record.need.organization
        return _sector_label((organization.sector if organization else None) or (record.data or {}).get("sector"))

    def title_of(record: BriefRecord) -> str:
        return (record.data or {}).get("title") or record.need.raw_text[:60]

    sectors: dict[str, Counter] = defaultdict(Counter)
    unmet: list[UnmetNeed] = []
    phrases: list[tuple[str, str]] = []
    funnel = Counter()
    for record in briefs:
        sector = sector_of(record)
        row = sectors[sector]
        row["needs"] += 1
        funnel["needs"] += 1
        if record.status != "final":
            continue
        funnel["briefed"] += 1
        run_id = latest_run.get(record.id)
        if run_id is not None:
            funnel["matched"] += 1
            if run_id not in shortlisted_runs:
                funnel["no_match"] += 1
                row["no_match"] += 1
                call = calls.get(record.id)
                unmet.append(
                    UnmetNeed(
                        brief_id=record.id,
                        title=title_of(record),
                        organization=record.need.organization.name if record.need.organization else None,
                        sector=sector,
                        required_capabilities=(record.data or {}).get("required_capabilities") or [],
                        open_call_id=call.id if call else None,
                        applications=app_counts.get(call.id, 0) if call else 0,
                    )
                )
                phrases += [(c, title_of(record)) for c in (record.data or {}).get("required_capabilities") or []]
            for rejection in session.scalars(
                select(Match.rejection).where(Match.run_id == run_id, Match.kind == "rejected")
            ):
                missing = (rejection or {}).get("missing_capability")
                if missing and missing.strip():
                    phrases.append((missing, title_of(record)))
        brief_pilots = pilots_by_brief.get(record.id, [])
        if record.id in intro_briefs or brief_pilots:
            funnel["introduced"] += 1
            row["introduced"] += 1
        if brief_pilots:
            funnel["piloted"] += 1
            row["pilots"] += len(brief_pilots)
        if any(p.result in ("evet", "kismen") for p in brief_pilots):
            funnel["worked"] += 1
            row["worked"] += 1

    answered = [i for i in intros if i.status in ("kabul", "ret") and i.responded_at]
    from_matching = [i for i in answered if i.application_id is None]
    accepted = sum(i.status == "kabul" for i in from_matching)
    intro_stats = IntroStats(
        total=len(intros),
        waiting=sum(i.status == "bekliyor" for i in intros),
        accepted=sum(i.status == "kabul" for i in intros),
        declined=sum(i.status == "ret" for i in intros),
        acceptance_rate=round(accepted / len(from_matching), 2) if from_matching else None,
        avg_response_days=round(
            sum((i.responded_at - i.created_at).total_seconds() for i in from_matching) / 86400 / len(from_matching), 1
        )
        if from_matching
        else None,
        via_admin=sum(i.responded_by == "yonetici" for i in intros),
    )
    pilot_stats = PilotStats(
        total=len(pilots),
        active=sum(p.status == "active" for p in pilots),
        stale=sum(p.status == "active" and (now - p.last_activity_at).days >= settings.pilot_stale_days for p in pilots),
        done=sum(p.status == "done" for p in pilots),
        result_evet=sum(p.result == "evet" for p in pilots),
        result_kismen=sum(p.result == "kismen" for p in pilots),
        result_hayir=sum(p.result == "hayir" for p in pilots),
    )

    owners = set(session.scalars(select(User.startup_id).where(User.startup_verified_at.is_not(None))))
    pool: dict[str, Counter] = defaultdict(Counter)
    for startup_id, sector in session.execute(select(Startup.id, Startup.sector).where(Startup.status == "aktif")):
        pool[_sector_label(sector)]["startups"] += 1
        pool[_sector_label(sector)]["with_account"] += startup_id in owners

    return ReportOut(
        generated_at=now,
        funnel=Funnel(**{k: funnel[k] for k in Funnel.model_fields}),
        introductions=intro_stats,
        pilots=pilot_stats,
        sectors=sorted(
            (SectorRow(sector=s, **{k: c[k] for k in ("needs", "no_match", "introduced", "pilots", "worked")}) for s, c in sectors.items()),
            key=lambda r: (-r.needs, r.sector),
        ),
        missing_capabilities=group_phrases(phrases, embedder)[:25],
        unmet_needs=unmet,
        pool=sorted((PoolRow(sector=s, **c) for s, c in pool.items()), key=lambda r: (-r.startups, r.sector)),
        open_calls=len(calls),
        applications=sum(app_counts.values()),
    )


@router.get("/report", response_model=ReportOut, summary="Program yöneticisi raporu")
def report(
    session: Session = Depends(get_db),
    _: User = Depends(require_admin),
    embedder: Embedder = Depends(embedder_dep),
    settings: Settings = Depends(get_settings),
):
    return build_report(session, embedder, settings)


def _pct(value: float | None) -> str:
    return "—" if value is None else f"%{round(value * 100)}"


@router.get("/report.xlsx", summary="Raporu Excel olarak indir")
def report_xlsx(
    session: Session = Depends(get_db),
    _: User = Depends(require_admin),
    embedder: Embedder = Depends(embedder_dep),
    settings: Settings = Depends(get_settings),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    data = build_report(session, embedder, settings)
    book = Workbook()

    def sheet(title: str, header: list[str], rows: list[list], first: bool = False):
        ws = book.active if first else book.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for i, column in enumerate(ws.columns, start=1):
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 70)
        ws.freeze_panes = "A2"

    f, i, p = data.funnel, data.introductions, data.pilots
    sheet(
        "Özet",
        ["Gösterge", "Değer"],
        [
            ["Rapor tarihi", data.generated_at.strftime("%d.%m.%Y %H:%M")],
            ["Girilen ihtiyaç", f.needs],
            ["Brief'i tamamlanan", f.briefed],
            ["Eşleştirmesi yapılan", f.matched],
            ["Havuzda doğrudan çözen girişim bulunamayan", f.no_match],
            ["Tanıştırma isteği giden ihtiyaç", f.introduced],
            ["Pilotu açılan ihtiyaç", f.piloted],
            ["Pilot sonucu olumlu (evet / kısmen)", f.worked],
            ["Tanıştırma isteği", i.total],
            ["Tanıştırma kabul oranı (eşleştirmeden)", _pct(i.acceptance_rate)],
            ["Ortalama cevap süresi (gün)", "—" if i.avg_response_days is None else i.avg_response_days],
            ["Yönetici aracılığıyla cevaplanan", i.via_admin],
            ["Pilot (toplam / aktif / hareketsiz / tamamlanan)", f"{p.total} / {p.active} / {p.stale} / {p.done}"],
            ["İşe yaradı mı (evet / kısmen / hayır)", f"{p.result_evet} / {p.result_kismen} / {p.result_hayir}"],
            ["Açık çağrı / başvuru", f"{data.open_calls} / {data.applications}"],
        ],
        first=True,
    )
    sheet(
        "Sektörler",
        ["Sektör", "İhtiyaç", "Uygun bulunamayan", "Tanıştırılan", "Pilot", "Olumlu sonuç"],
        [[r.sector, r.needs, r.no_match, r.introduced, r.pilots, r.worked] for r in data.sectors],
    )
    sheet(
        "Eksik yetkinlikler",
        ["Yetkinlik", "İhtiyaç sayısı", "Diğer yazımlar", "İhtiyaçlar"],
        [[m.capability, m.needs, "; ".join(m.variants), "; ".join(m.need_titles)] for m in data.missing_capabilities],
    )
    sheet(
        "Uygun bulunamayan",
        ["İhtiyaç", "Kurum", "Sektör", "Aranan yetkinlikler", "Açık çağrı", "Başvuru"],
        [
            [u.title, u.organization or "", u.sector, "; ".join(u.required_capabilities), "var" if u.open_call_id else "yok", u.applications]
            for u in data.unmet_needs
        ],
    )

    pilot_rows = []
    for pilot in session.scalars(select(Pilot).order_by(Pilot.id)):
        record = session.get(BriefRecord, pilot.brief_id)
        pilot_rows.append([
            (record.data or {}).get("title") or record.need.raw_text[:60],
            record.need.organization.name if record.need.organization else "",
            session.get(Startup, pilot.startup_id).name,
            pilot.status,
            pilot.result or "",
            pilot.started_at.strftime("%d.%m.%Y"),
            pilot.last_activity_at.strftime("%d.%m.%Y"),
            pilot.outcome or "",
        ])
    sheet("Pilotlar", ["İhtiyaç", "Kurum", "Girişim", "Durum", "İşe yaradı mı", "Başlangıç", "Son hareket", "Not"], pilot_rows)
    sheet("Havuz", ["Sektör", "Girişim", "Hesabı olan"], [[r.sector, r.startups, r.with_account] for r in data.pool])

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    filename = f"needle-rapor-{data.generated_at:%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
