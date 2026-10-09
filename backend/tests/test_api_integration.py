"""Uçtan uca API testi: gerçek Postgres + pgvector, sahte LLM.

Çalıştırmak için (docker compose up -d db sonrası):
    NEEDLE_TEST_DATABASE_URL=postgresql+psycopg://needle:needle@127.0.0.1:5432/needle python -m pytest tests/test_api_integration.py
"""

import os

import pytest

DB_URL = os.getenv("NEEDLE_TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not DB_URL, reason="Entegrasyon testi: NEEDLE_TEST_DATABASE_URL gerekli")


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", DB_URL)
    from fastapi.testclient import TestClient
    from sqlalchemy import text

    from app.api import deps
    from app.config import Settings, get_settings
    from app.db import session as db_session
    from app.db.models import Base, Startup
    from app.embeddings import HashingEmbedder
    from app.llm.client import StructuredLLM
    from app.main import app
    from app.rerank.rerankers import PassthroughReranker
    from app.seed_data import load_startups
    from tests.test_brief_graph import COMPLETE_BRIEF, PARTIAL_BRIEF, QUESTIONS
    from tests.test_match_graph import EchoRationaleBackend

    get_settings.cache_clear()
    db_session.get_engine.cache_clear()
    engine = db_session.get_engine()
    Base.metadata.drop_all(engine)
    db_session.init_db(engine)

    embedder = HashingEmbedder()
    with db_session.get_session_factory()() as s:
        for p, vector in zip(load_startups(), embedder.embed_documents([p.to_search_text() for p in load_startups()])):
            s.add(Startup(
                id=p.id, name=p.name, sector=p.sector, maturity=p.maturity.value, location=p.location,
                capabilities=p.capabilities, description=p.description, past_pilots=p.past_pilots, embedding=vector,
            ))
        s.commit()

    settings = Settings(database_url=DB_URL, demo_cache_dir=str(tmp_path))

    class RoutingBackend:
        """Brief çağrılarına sırayla yazılı cevap, gerekçe çağrısına EchoRationaleBackend."""

        model_name = "fake-model"

        def __init__(self):
            self.brief_responses = [PARTIAL_BRIEF, QUESTIONS, COMPLETE_BRIEF]
            self.rationale = EchoRationaleBackend()

        def complete(self, messages):
            if "Kısa liste (matches)" in messages[1]["content"]:
                return self.rationale.complete(messages)
            return self.brief_responses.pop(0)

    llm = StructuredLLM(RoutingBackend(), settings)
    app.dependency_overrides[deps.llm_dep] = lambda: llm
    app.dependency_overrides[deps.embedder_dep] = lambda: embedder
    app.dependency_overrides[deps.reranker_dep] = lambda: PassthroughReranker()
    app.dependency_overrides[get_settings] = lambda: settings
    yield TestClient(app)
    app.dependency_overrides.clear()
    with engine.begin() as conn:
        conn.execute(text("SELECT 1"))


def register(client, email: str, organization: str | None = None, account_type: str = "firma"):
    response = client.post("/auth/register", json={
        "name": "Deneme Kullanıcı", "email": email, "password": "uzun-bir-sifre-123",
        "organization_name": organization, "account_type": account_type, "kvkk_onay": True,
    })
    assert response.status_code == 200, response.text
    return response.json()


def admin_client():
    from fastapi.testclient import TestClient

    from app.auth import hash_password
    from app.db import session as db_session
    from app.db.models import User
    from app.main import app

    with db_session.get_session_factory()() as s:
        if not s.query(User).filter_by(email="yonetici@program.example").first():
            s.add(User(email="yonetici@program.example", password_hash=hash_password("yonetici-sifresi-1"),
                       name="Program Yöneticisi", role="yonetici"))
            s.commit()
    admin = TestClient(app)
    assert admin.post("/auth/login", json={"email": "yonetici@program.example", "password": "yonetici-sifresi-1"}).status_code == 200
    return admin


def matched_need(client) -> tuple[int, dict]:
    """Firma oturumunda tamamlanmış brief + eşleştirme sonucu."""
    body = client.post("/needs", json={
        "raw_text": "Sahadaki bayilerimizden gelen şikayetleri daha hızlı sınıflandırmak istiyoruz, çok manuel gidiyor.",
    }).json()
    client.post(f"/briefs/{body['brief_id']}/answers", json={"answers": {
        "scope": "Ayda 3.000 şikayet", "success_criteria": "%85 doğruluk", "timeline": "3 ay",
    }})
    return body["brief_id"], client.post(f"/briefs/{body['brief_id']}/match").json()


def insert_brief(email: str, capabilities: list[str]) -> int:
    from app.db import session as db_session
    from app.db.models import BriefRecord, Need, User

    with db_session.get_session_factory()() as s:
        org_id = s.query(User).filter_by(email=email).one().organization_id
        need = Need(organization_id=org_id, raw_text="Depolarda stok sayımı ve şikayet ayrıştırma")
        s.add(need)
        s.flush()
        record = BriefRecord(need_id=need.id, status="final", followup_questions=[], answers={}, data={
            "title": "Depo ve şikayet ihtiyacı", "problem": "Elle yapılıyor", "required_capabilities": capabilities,
        })
        s.add(record)
        s.commit()
        return record.id


def set_website(startup_id: str, url: str):
    from app.db import session as db_session
    from app.db.models import Startup

    with db_session.get_session_factory()() as s:
        s.get(Startup, startup_id).website = url
        s.commit()


def test_full_flow_need_to_pilot(client):
    assert client.get("/needs").status_code == 401
    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")

    need = client.post("/needs", json={
        "raw_text": "Sahadaki bayilerimizden gelen şikayetleri daha hızlı sınıflandırmak istiyoruz, çok manuel gidiyor.",
        "organization": {"name": "Kuzey Beyaz Eşya", "author_unit": "İnovasyon", "owner_unit": "Bayi Satış"},
    })
    assert need.status_code == 200, need.text
    body = need.json()
    assert body["status"] == "needs_input"
    assert {q["field"] for q in body["questions"]} == {"scope", "success_criteria", "timeline"}

    assert client.post(f"/briefs/{body['brief_id']}/match").status_code == 409

    answered = client.post(f"/briefs/{body['brief_id']}/answers", json={"answers": {
        "scope": "Ayda 3.000 şikayet", "success_criteria": "%85 doğruluk", "timeline": "3 ay",
    }})
    assert answered.status_code == 200, answered.text
    assert answered.json()["status"] == "final"

    match = client.post(f"/briefs/{body['brief_id']}/match")
    assert match.status_code == 200, match.text
    result = match.json()
    assert len(result["shortlist"]) == 5 and len(result["rejected"]) == 3
    assert result["shortlist"][0]["startup"]["id"] == "s01"
    assert result["retrieval_trace"]

    top_match_id = result["match_ids"]["s01"]
    decision = client.post(f"/matches/{top_match_id}/decision", json={"decision": "accept", "note": "Görüşelim"})
    assert decision.status_code == 200
    intro_id = decision.json()["introduction_id"]
    assert client.get("/pilots").json() == []
    [intro] = client.get("/introductions").json()
    assert intro["status"] == "bekliyor" and intro["firm_note"] == "Görüşelim" and not intro["startup_has_account"]
    assert client.post(f"/introductions/{intro_id}/respond", json={"decision": "kabul"}).status_code == 403
    answered = admin_client().post(f"/introductions/{intro_id}/respond", json={"decision": "kabul", "note": "Telefonla görüşüldü"})
    assert answered.status_code == 200 and answered.json()["responded_by"] == "yonetici"

    again = client.post(f"/matches/{top_match_id}/decision", json={"decision": "decline"})
    assert again.status_code == 409

    declined_id = result["match_ids"][result["shortlist"][1]["startup"]["id"]]
    client.post(f"/matches/{declined_id}/decision", json={"decision": "decline", "reason": "Bütçe uymadı"})

    needs = client.get("/needs").json()
    assert needs[0]["brief_id"] == body["brief_id"] and needs[0]["accepted_count"] == 1
    assert needs[0]["organization"] == "Kuzey Beyaz Eşya"

    saved = client.get(f"/briefs/{body['brief_id']}/match").json()
    by_id = {item["match_id"]: item for item in saved["shortlist"]}
    assert by_id[top_match_id]["status"] == "accepted"
    assert by_id[top_match_id]["introduction"]["status"] == "kabul"
    assert by_id[declined_id]["declined_reason"] == "Bütçe uymadı" and by_id[declined_id]["rejection"] is None

    pilot_id = answered.json()["pilot_id"]
    pilot = client.post(f"/pilots/{pilot_id}/milestones", json={"title": "Veri erişimi", "due_date": "2026-11-01"}).json()
    assert pilot["startup"]["id"] == "s01" and pilot["stale"] is False
    pilot = client.post(f"/milestones/{pilot['milestones'][0]['id']}/complete").json()
    assert pilot["milestones"][0]["completed_at"]
    pilot = client.patch(f"/pilots/{pilot_id}", json={"status": "done", "result": "kismen"}).json()
    assert pilot["status"] == "done" and pilot["result"] == "kismen"
    assert [p["id"] for p in client.get("/pilots").json()] == [pilot_id]

    assert len(client.get("/startups").json()) == 40


def test_firms_only_see_their_own_records(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "a@firma-a.example", "Firma A")
    created = client.post("/needs", json={"raw_text": "Sahadaki bayilerimizden gelen şikayetleri sınıflandırmak istiyoruz."})
    brief_id = created.json()["brief_id"]

    other = TestClient(app)
    register(other, "b@firma-b.example", "Firma B")
    assert other.get("/needs").json() == []
    assert other.get(f"/briefs/{brief_id}").status_code == 404
    assert other.post(f"/briefs/{brief_id}/answers", json={"answers": {"scope": "x"}}).status_code == 404
    assert other.post(f"/briefs/{brief_id}/match").status_code == 404
    assert other.get("/pilots").json() == []

    admin = admin_client()
    assert TestClient(app).post("/auth/login", json={"email": "yonetici@program.example", "password": "yanlis"}).status_code == 401
    assert TestClient(app).post("/auth/login", json={"email": "YONETICI@program.example", "password": "yonetici-sifresi-1"}).status_code == 200
    assert [n["brief_id"] for n in admin.get("/needs").json()] == [brief_id]
    assert admin.get(f"/briefs/{brief_id}").status_code == 200

    assert other.post("/auth/logout").status_code == 204
    assert other.get("/needs").status_code == 401


def test_register_rejects_duplicates_and_missing_consent(client):
    register(client, "tek@firma.example", "Firma")
    duplicate = client.post("/auth/register", json={
        "name": "Başka", "email": "TEK@firma.example", "password": "uzun-bir-sifre-123",
        "organization_name": "Başka Firma", "kvkk_onay": True,
    })
    assert duplicate.status_code == 409
    no_consent = client.post("/auth/register", json={
        "name": "Başka", "email": "yeni@firma.example", "password": "uzun-bir-sifre-123",
        "organization_name": "Başka Firma", "kvkk_onay": False,
    })
    assert no_consent.status_code == 422


def test_profile_fills_empty_brief_fields(client):
    register(client, "profil@firma.example", "Profil Firma")
    me = client.put("/auth/profile", json={
        "sector": "perakende", "city": "İstanbul", "employee_range": "250-999", "systems": ["SAP"],
        "preferred_maturity": "mvp", "startup_location": "ayni_sehir", "pilot_duration": "3 ay",
    }).json()
    assert me["organization"]["onboarded"] is True

    body = client.post("/needs", json={"raw_text": "Sahadaki bayilerimizden gelen şikayetleri sınıflandırmak istiyoruz."}).json()
    brief = body["brief"]
    assert brief["sector"] == "perakende"
    assert brief["location_preference"] == "İstanbul"
    assert brief["min_maturity"] == "mvp"
    assert brief["timeline"] == "3 ay"
    assert "timeline" not in {q["field"] for q in body["questions"]}


def test_startup_claims_profile_and_answers_introduction(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")
    brief_id, result = matched_need(client)
    intro_id = client.post(f"/matches/{result['match_ids']['s01']}/decision", json={"decision": "accept"}).json()["introduction_id"]

    set_website("s01", "https://www.metinsel.example/tr")
    startup = TestClient(app)
    me = register(startup, "ali@metinsel.example", account_type="girisim")
    assert me["role"] == "girisim" and me["organization"] is None and me["startup"] is None

    assert startup.get("/needs").status_code == 403
    assert startup.get(f"/briefs/{brief_id}").status_code == 403
    assert startup.get(f"/briefs/{brief_id}/match").status_code == 403
    assert startup.get("/introductions").status_code == 403

    me = startup.post("/startup-account/claim", json={"startup_id": "s01"}).json()
    assert me["startup"]["verified"] is False
    admin = admin_client()
    [claim] = admin.get("/admin/claims").json()
    assert claim["domain_match"] is True
    assert admin.post(f"/admin/claims/{claim['user_id']}/approve").status_code == 204
    assert startup.get("/auth/me").json()["startup"]["verified"] is True

    [intro] = startup.get("/introductions").json()
    assert intro["id"] == intro_id and intro["organization"] == "Kuzey Beyaz Eşya"
    assert intro["brief"]["required_capabilities"]
    answered = startup.post(f"/introductions/{intro_id}/respond", json={"decision": "kabul", "note": "Memnuniyetle"}).json()
    assert answered["status"] == "kabul" and answered["responded_by"] == "girisim"
    assert startup.post(f"/introductions/{intro_id}/respond", json={"decision": "ret"}).status_code == 409

    [pilot] = startup.get("/pilots").json()
    assert pilot["id"] == answered["pilot_id"] and [p["id"] for p in client.get("/pilots").json()] == [pilot["id"]]
    assert startup.post(f"/pilots/{pilot['id']}/milestones", json={"title": "Örnek veri"}).status_code == 200
    assert startup.patch(f"/pilots/{pilot['id']}", json={"result": "evet"}).status_code == 403

    call_id = client.post("/calls", json={
        "brief_id": brief_id, "title": "Şikayet sınıflandırma çağrısı",
        "summary": "Ayda 3.000 Türkçe şikayet metnini otomatik sınıflandırmak istiyoruz.",
    }).json()["id"]
    note = {"note": "Hazır modelimiz var, iki haftada kurarız, referanslarımız mevcut."}
    assert startup.post(f"/calls/{call_id}/applications", json=note).status_code == 409

    other = TestClient(app)
    register(other, "veli@metinsel.example", account_type="girisim")
    assert other.post("/startup-account/claim", json={"startup_id": "s01"}).status_code == 409


def test_unverified_claim_and_new_profile_need_admin_approval(client):
    register(client, "kisisel@gmail.com", account_type="girisim")
    set_website("s02", "https://s02.example")
    me = client.post("/startup-account/claim", json={"startup_id": "s02"}).json()
    assert me["startup"]["verified"] is False

    from fastapi.testclient import TestClient

    from app.main import app

    newcomer = TestClient(app)
    register(newcomer, "kurucu@yeni-girisim.example", account_type="girisim")
    profile = {
        "name": "Yeni Girişim", "sector": "enerji", "maturity": "mvp", "location": "İzmir",
        "website": "https://yeni-girisim.example", "description": "Akustik sensörle su şebekesinde kaçak tespiti yapıyoruz.",
        "capabilities": ["akustik sensörle su kaçağı tespiti", "şebeke basınç izleme"],
    }
    me = newcomer.post("/startup-account/new", json=profile).json()
    assert me["startup"]["status"] == "onay_bekliyor"
    new_id = me["startup"]["id"]
    assert new_id not in {s["id"] for s in client.get("/startups").json()}

    admin = admin_client()
    claims = {c["kind"]: c for c in admin.get("/admin/claims").json()}
    assert claims["sahiplenme"]["domain_match"] is False and claims["yeni_profil"]["startup"]["id"] == new_id
    assert admin.post(f"/admin/claims/{claims['yeni_profil']['user_id']}/approve").status_code == 204
    assert admin.post(f"/admin/claims/{claims['sahiplenme']['user_id']}/reject").status_code == 204

    assert newcomer.get("/auth/me").json()["startup"] == {"id": new_id, "name": "Yeni Girişim", "status": "aktif", "verified": True}
    assert new_id in {s["id"] for s in admin.get("/startups").json()}
    assert client.get("/auth/me").json()["startup"] is None
    assert client.get("/startup-account/profile").status_code == 404


def test_open_call_application_becomes_pilot(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")
    brief_id, _ = matched_need(client)
    call = client.post("/calls", json={
        "brief_id": brief_id, "title": "Bayi şikayetlerini sınıflandıracak girişim arıyoruz",
        "summary": "Ayda 3.000 Türkçe şikayet metnini %85 doğrulukla otomatik sınıflandırmak istiyoruz.",
        "hide_organization": True,
    })
    assert call.status_code == 200, call.text
    call_id = call.json()["id"]
    assert client.post("/calls", json={**call.json(), "brief_id": brief_id}).status_code == 409
    assert client.get(f"/briefs/{brief_id}/match").json()["open_call_id"] == call_id

    set_website("s03", "https://s03.example")
    startup = TestClient(app)
    register(startup, "ekip@s03.example", account_type="girisim")
    startup.post("/startup-account/claim", json={"startup_id": "s03"})
    [claim] = admin_client().get("/admin/claims").json()
    admin_client().post(f"/admin/claims/{claim['user_id']}/approve")
    [listed] = startup.get("/calls").json()
    assert listed["organization"] is None and listed["application_count"] == 0
    note = {"note": "Türkçe şikayet metinlerini sınıflandıran hazır modelimiz var, iki haftada kurarız."}
    applied = startup.post(f"/calls/{call_id}/applications", json=note).json()
    assert applied["my_application"]["status"] == "yeni"
    assert startup.post(f"/calls/{call_id}/applications", json=note).status_code == 409

    other = TestClient(app)
    register(other, "b@firma-b.example", "Firma B")
    assert other.get("/calls").json() == [] and other.get(f"/calls/{call_id}").status_code == 404

    detail = client.get(f"/calls/{call_id}").json()
    assert detail["organization"] == "Kuzey Beyaz Eşya" and detail["application_count"] == 1
    [app_] = detail["applications"]
    assert startup.post(f"/applications/{app_['id']}/decision", json={"decision": "kabul"}).status_code == 403
    decided = client.post(f"/applications/{app_['id']}/decision", json={"decision": "kabul", "note": "Başlayalım"}).json()
    assert decided["applications"][0]["status"] == "kabul" and decided["applications"][0]["pilot_id"]

    [pilot] = startup.get("/pilots").json()
    assert pilot["startup"]["id"] == "s03" and pilot["match_id"] is None
    [intro] = startup.get("/introductions").json()
    assert intro["source"] == "cagri" and intro["status"] == "kabul"

    assert client.patch(f"/calls/{call_id}", json={"status": "kapali"}).json()["status"] == "kapali"
    assert [c["id"] for c in startup.get("/calls").json()] == [call_id]


def test_domain_autoverify_only_when_enabled(client):
    from app.api import collab_routes  # noqa: F401  (ayarın uç noktaya ulaştığını sınar)
    from app.config import get_settings
    from app.main import app

    set_website("s04", "https://s04.example")
    enabled = get_settings().model_copy(update={"startup_domain_autoverify": True})
    app.dependency_overrides[get_settings] = lambda: enabled
    register(client, "kurucu@s04.example", account_type="girisim")
    assert client.post("/startup-account/claim", json={"startup_id": "s04"}).json()["startup"]["verified"] is True


def test_admin_report_and_excel_export(client):
    from io import BytesIO

    from openpyxl import load_workbook

    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")
    client.put("/auth/profile", json={"sector": "Perakende", "city": "İstanbul", "employee_range": "250-999"})
    brief_id, result = matched_need(client)
    intro_id = client.post(f"/matches/{result['match_ids']['s01']}/decision", json={"decision": "accept"}).json()["introduction_id"]
    admin = admin_client()
    pilot_id = admin.post(f"/introductions/{intro_id}/respond", json={"decision": "kabul"}).json()["pilot_id"]
    client.patch(f"/pilots/{pilot_id}", json={"status": "done", "result": "evet"})

    assert client.get("/admin/report").status_code == 403
    report = admin.get("/admin/report").json()
    assert report["funnel"] == {
        "needs": 1, "briefed": 1, "matched": 1, "no_match": 0, "introduced": 1, "piloted": 1, "worked": 1,
    }
    assert report["introductions"]["acceptance_rate"] == 1.0 and report["introductions"]["via_admin"] == 1
    assert report["pilots"]["result_evet"] == 1
    assert report["sectors"] == [{"sector": "Perakende", "needs": 1, "no_match": 0, "introduced": 1, "pilots": 1, "worked": 1}]
    assert sum(r["startups"] for r in report["pool"]) == 40

    excel = admin.get("/admin/report.xlsx")
    assert excel.status_code == 200 and "needle-rapor-" in excel.headers["content-disposition"]
    book = load_workbook(BytesIO(excel.content))
    assert book.sheetnames == ["Özet", "Sektörler", "Eksik yetkinlikler", "Uygun bulunamayan", "Pilotlar", "Havuz"]
    assert book["Pilotlar"]["E2"].value == "evet"


def test_login_is_locked_after_repeated_failures(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "kilit@firma.example", "Kilit Firma")
    attacker = TestClient(app)
    for _ in range(5):
        assert attacker.post("/auth/login", json={"email": "kilit@firma.example", "password": "yanlis-sifre"}).status_code == 401
    locked = attacker.post("/auth/login", json={"email": "KILIT@firma.example", "password": "uzun-bir-sifre-123"})
    assert locked.status_code == 429 and "dakika" in locked.json()["detail"]
    register(TestClient(app), "baska@firma.example", "Başka Firma")
    assert TestClient(app).post("/auth/login", json={"email": "baska@firma.example", "password": "uzun-bir-sifre-123"}).status_code == 200


def test_password_reset_flow(client, monkeypatch):
    from fastapi.testclient import TestClient

    from app.api import auth_routes
    from app.main import app

    sent = []
    monkeypatch.setattr(auth_routes, "send_mail", lambda settings, to, subject, body: sent.append((to, body)))
    register(client, "unuttum@firma.example", "Unutkan Firma")
    for _ in range(5):
        TestClient(app).post("/auth/login", json={"email": "unuttum@firma.example", "password": "yanlis-sifre"})

    assert client.post("/auth/password-reset/request", json={"email": "yok@firma.example"}).status_code == 204
    assert client.post("/auth/password-reset/request", json={"email": "Unuttum@firma.example"}).status_code == 204
    [(to, body)] = sent
    token = body.split("anahtar=")[1].split()[0]
    assert to == "unuttum@firma.example"

    assert client.post("/auth/password-reset/confirm", json={"token": token, "password": "kisa"}).status_code == 422
    assert client.post("/auth/password-reset/confirm", json={"token": token, "password": "yepyeni-sifre-456"}).status_code == 204
    assert client.get("/auth/me").status_code == 401
    assert client.post("/auth/password-reset/confirm", json={"token": token, "password": "baska-sifre-789"}).status_code == 400
    assert client.post("/auth/login", json={"email": "unuttum@firma.example", "password": "uzun-bir-sifre-123"}).status_code == 401
    assert client.post("/auth/login", json={"email": "unuttum@firma.example", "password": "yepyeni-sifre-456"}).status_code == 200


def test_change_password_signs_out_other_devices(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "degis@firma.example", "Değişen Firma")
    other_device = TestClient(app)
    other_device.post("/auth/login", json={"email": "degis@firma.example", "password": "uzun-bir-sifre-123"})
    assert other_device.get("/auth/me").status_code == 200

    wrong = client.post("/auth/password", json={"current_password": "yanlis", "new_password": "yepyeni-sifre-456"})
    assert wrong.status_code == 400
    assert client.post("/auth/password", json={"current_password": "uzun-bir-sifre-123", "new_password": "yepyeni-sifre-456"}).status_code == 204
    assert client.get("/auth/me").status_code == 200
    assert other_device.get("/auth/me").status_code == 401


def test_intro_email_draft_for_startups_without_account(client):
    from fastapi.testclient import TestClient

    from app.db import session as db_session
    from app.db.models import Startup
    from app.main import app

    with db_session.get_session_factory()() as s:
        s.get(Startup, "s01").contact_email = "info@metinsel.example"
        s.get(Startup, "s01").contact_source = "https://metinsel.example/iletisim"
        s.commit()
    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")
    _, result = matched_need(client)
    intro_id = client.post(
        f"/matches/{result['match_ids']['s01']}/decision", json={"decision": "accept", "note": "Haftaya görüşebilir miyiz?"}
    ).json()["introduction_id"]

    [intro] = client.get("/introductions").json()
    email = intro["email"]
    assert email["to"] == "info@metinsel.example" and email["sent_at"] is None
    assert "Kuzey Beyaz Eşya sizinle tanışmak istiyor" in email["subject"]
    assert "uygun." in email["body"] and "Haftaya görüşebilir miyiz?" in email["body"] and "/kayit" in email["body"]

    edited = client.put(f"/introductions/{intro_id}/email", json={**email, "subject": "Kısa bir tanışma", "to": "ceo@metinsel.example"})
    assert edited.status_code == 200 and edited.json()["email"]["subject"] == "Kısa bir tanışma"
    assert client.put(f"/introductions/{intro_id}/email", json={**email, "to": "adres-degil"}).status_code == 422
    sent = client.post(f"/introductions/{intro_id}/email/sent", json={"sent": True}).json()["email"]
    assert sent["sent_at"] is not None
    assert client.put(f"/introductions/{intro_id}/email", json=email).status_code == 409

    other = TestClient(app)
    register(other, "b@firma-b.example", "Firma B")
    assert other.put(f"/introductions/{intro_id}/email", json=email).status_code == 404

    startup = TestClient(app)
    register(startup, "ali@metinsel.example", account_type="girisim")
    startup.post("/startup-account/claim", json={"startup_id": "s01"})
    [claim] = admin_client().get("/admin/claims").json()
    admin_client().post(f"/admin/claims/{claim['user_id']}/approve")
    [seen] = startup.get("/introductions").json()
    assert seen["email"] is None
    assert startup.post(f"/introductions/{intro_id}/email/sent", json={"sent": False}).status_code == 404


def test_startups_can_explore_organizations_before_verification(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")
    client.put("/auth/profile", json={
        "sector": "Perakende", "city": "İstanbul", "employee_range": "250-999", "systems": ["SAP"],
        "budget_range": "1 milyon TL üzeri", "description": "Beyaz eşya üretip bayi ağıyla satıyoruz.",
        "website": "https://kuzey.example",
    })
    brief_id, _ = matched_need(client)
    visible = client.post("/calls", json={
        "brief_id": brief_id, "title": "Şikayet sınıflandırma çağrısı", "summary": "Ayda 3.000 şikayeti otomatik sınıflandırmak istiyoruz.",
    }).json()["id"]

    gizli = TestClient(app)
    register(gizli, "gizli@firma.example", "Gizli Firma")
    gizli.put("/auth/profile", json={"sector": "Enerji", "city": "Ankara", "employee_range": "1000+", "directory_visible": False})
    register(TestClient(app), "yarim@firma.example", "Profili Eksik Firma")

    startup = TestClient(app)
    register(startup, "kurucu@yeni.example", account_type="girisim")
    [org] = startup.get("/organizations").json()
    assert org["name"] == "Kuzey Beyaz Eşya" and org["description"].startswith("Beyaz eşya")
    assert org["city"] == "İstanbul" and org["website"] == "https://kuzey.example"
    assert [c["id"] for c in org["open_calls"]] == [visible]
    assert "systems" not in org and "budget_range" not in org
    assert len(startup.get("/startups").json()) == 40
    assert [c["id"] for c in startup.get("/calls").json()] == [visible]
    note = {"note": "Hazır modelimiz var, iki haftada kurarız, referanslarımız mevcut."}
    assert startup.post(f"/calls/{visible}/applications", json=note).status_code == 403
    assert startup.get("/introductions").status_code == 403

    client.patch(f"/calls/{visible}", json={"status": "kapali"})
    assert startup.get("/organizations").json()[0]["open_calls"] == []


def test_pilot_management_plan_metrics_activity_and_evaluation(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "inovasyon@kuzey.example", "Kuzey Beyaz Eşya")
    brief_id, result = matched_need(client)
    intro_id = client.post(f"/matches/{result['match_ids']['s01']}/decision", json={"decision": "accept"}).json()["introduction_id"]
    pilot_id = admin_client().post(f"/introductions/{intro_id}/respond", json={"decision": "kabul"}).json()["pilot_id"]

    # Pilot açılınca plan brief'ten önerilir: amaç, tarihler (3 ay), 5 kilometre taşı ve başarı kriterinden hedef
    pilot = client.get(f"/pilots/{pilot_id}").json()
    assert pilot["goal"] and pilot["start_date"] and pilot["end_date"]
    assert [m["owner"] for m in pilot["milestones"]] == ["ortak", "kurum", "girisim", "ortak", "ortak"]
    [metric] = pilot["metrics"]
    assert metric["target"] == 85.0 and metric["unit"] == "%"
    assert "Pilot açıldı" in pilot["activity"][-1]["body"]

    pilot = client.patch(f"/pilots/{pilot_id}/plan", json={"firm_contact": "Ayşe Yılmaz, ayse@kuzey.example"}).json()
    assert pilot["firm_contact"].startswith("Ayşe")
    assert client.patch(f"/pilots/{pilot_id}/plan", json={"start_date": "2026-12-01", "end_date": "2026-11-01"}).status_code == 422

    client.patch(f"/metrics/{metric['id']}", json={**{k: metric[k] for k in ("name", "unit", "target", "direction")}, "baseline": 60})
    pilot = client.post(f"/metrics/{metric['id']}/measurements", json={"value": 72.5, "measured_on": "2026-11-15"}).json()
    assert pilot["metrics"][0]["latest"] == 72.5 and pilot["metrics"][0]["progress"] == 0.5

    late = client.post(f"/pilots/{pilot_id}/milestones", json={"title": "Geciken iş", "due_date": "2020-01-01", "owner": "girisim"}).json()
    assert late["overdue_milestones"] == 1

    # Girişim pilotu görür, not yazar, ölçüm ekler; durumu ve kurumun değerlendirmesini değiştiremez
    startup = TestClient(app)
    register(startup, "ali@metinsel.example", account_type="girisim")
    set_website("s01", "https://metinsel.example")
    startup.post("/startup-account/claim", json={"startup_id": "s01"})
    [claim] = admin_client().get("/admin/claims").json()
    admin_client().post(f"/admin/claims/{claim['user_id']}/approve")
    assert startup.post(f"/pilots/{pilot_id}/activity", json={"body": "Örnek veri setini aldık, kuruluma başlıyoruz."}).status_code == 200
    assert startup.post(f"/metrics/{metric['id']}/measurements", json={"value": 80}).status_code == 200
    assert startup.patch(f"/pilots/{pilot_id}", json={"status": "paused"}).status_code == 403
    evaluation = {"result": "evet", "next_step": "satin_alma", "startup_rating": 5, "comment": "Hedef tuttu."}
    assert startup.post(f"/pilots/{pilot_id}/evaluation", json=evaluation).status_code == 403
    assert client.post(f"/pilots/{pilot_id}/startup-feedback", json={"feedback": "Kurum yazmamalı bunu."}).status_code == 403

    # Başka firma göremez
    other = TestClient(app)
    register(other, "b@firma-b.example", "Firma B")
    assert other.get(f"/pilots/{pilot_id}").status_code == 404
    assert other.post(f"/pilots/{pilot_id}/activity", json={"body": "merhaba"}).status_code == 404

    done = client.post(f"/pilots/{pilot_id}/evaluation", json=evaluation).json()
    assert done["status"] == "done" and done["next_step"] == "satin_alma" and done["startup_rating"] == 5
    fb = startup.post(f"/pilots/{pilot_id}/startup-feedback", json={"feedback": "Veri erişimi hızlıydı, iyi bir iş birliği oldu.", "collab_rating": 4}).json()
    assert fb["collab_rating"] == 4 and fb["startup_feedback_at"]

    bodies = [a["body"] for a in client.get(f"/pilots/{pilot_id}").json()["activity"]]
    assert bodies[0] == "Girişim pilotu değerlendirdi (iş birliği 4/5)"
    assert any(b.startswith("Pilot değerlendirildi: işe yaradı mı evet") for b in bodies)
    assert "Örnek veri setini aldık, kuruluma başlıyoruz." in bodies
    assert any(b.startswith("Ölçüm: ") and b.endswith(" = 72,5%") for b in bodies)


def test_admin_account_is_bootstrapped_from_settings(client):
    from seed.create_admin import ensure_admin

    assert ensure_admin("Kurulum@Program.example", "kisa", "Ad") .startswith("ADMIN_PASSWORD en az 10")
    assert ensure_admin("Kurulum@Program.example", "kurulum-sifresi-1", "Kurulum Yöneticisi") == "Yönetici hesabı oluşturuldu: kurulum@program.example"
    assert ensure_admin("kurulum@program.example", "kurulum-sifresi-1", "X") == "Yönetici hesabı zaten var: kurulum@program.example"
    login = client.post("/auth/login", json={"email": "kurulum@program.example", "password": "kurulum-sifresi-1"})
    assert login.status_code == 200 and login.json()["role"] == "yonetici"
    register(client, "firma@ornek.example", "Örnek Firma")
    assert "yönetici yapılmadı" in ensure_admin("firma@ornek.example", "kurulum-sifresi-1", "X")


def test_startup_recommendations_calls_organizations_and_demand_signals(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "a@firma-a.example", "Firma A")
    brief_a, _ = matched_need(client)
    client.post("/calls", json={"brief_id": brief_a, "title": "Türkçe metin sınıflandırma çağrısı",
                                "summary": "Bayi şikayetlerini Türkçe metin sınıflandırma ile ayırmak istiyoruz."})
    other = TestClient(app)
    register(other, "b@firma-b.example", "Firma B")
    other.put("/auth/profile", json={"sector": "Lojistik", "city": "İzmir", "employee_range": "250-999",
                                     "description": "Depo ve stok yönetimi yapan lojistik firması."})
    brief_b = insert_brief("b@firma-b.example", ["Türkçe metin sınıflandırma", "RFID ile stok sayımı"])
    other.post("/calls", json={"brief_id": brief_b, "title": "Depo stok sayımı çağrısı",
                               "summary": "Depolarda RFID ile stok sayımını hızlandıracak çözüm arıyoruz."})

    startup = TestClient(app)
    register(startup, "ali@metinsel.example", account_type="girisim")
    assert startup.get("/startup-account/recommendations").status_code == 409
    startup.post("/startup-account/claim", json={"startup_id": "s01"})
    recs = startup.get("/startup-account/recommendations").json()

    assert recs["calls"][0]["title"] == "Türkçe metin sınıflandırma çağrısı"
    assert "Türkçe metin sınıflandırma" in recs["calls"][0]["reason"]
    assert recs["calls"][0]["score"] > recs["calls"][1]["score"]
    assert [(s["capability"], s["organizations"]) for s in recs["signals"]] == [("Türkçe metin sınıflandırma", 2)]
    assert set(recs["signals"][0]) == {"capability", "organizations", "variants", "score"}
    assert "Firma B" in {o["name"] for o in recs["organizations"]}
    assert client.get("/startup-account/recommendations").status_code == 403


def test_organization_recommendations_mirror_startup_side(client):
    from fastapi.testclient import TestClient

    from app.main import app

    register(client, "a@firma-a.example", "Firma A")
    assert client.get("/organization/recommendations").status_code == 409
    brief_id, _ = matched_need(client)
    call = client.post("/calls", json={"brief_id": brief_id, "title": "Şikayet sınıflandırma çağrısı",
                                       "summary": "Bayi şikayetlerini otomatik ayırmak istiyoruz."}).json()

    startup = TestClient(app)
    register(startup, "ali@metinsel.example", account_type="girisim")
    startup.post("/startup-account/claim", json={"startup_id": "s01"})
    admin = admin_client()
    [claim] = admin.get("/admin/claims").json()
    admin.post(f"/admin/claims/{claim['user_id']}/approve")
    assert startup.post(f"/calls/{call['id']}/applications", json={"note": "Türkçe şikayet metinlerini sınıflandıran hazır modelimiz var."}).status_code == 200
    assert startup.get("/organization/recommendations").status_code == 403

    recs = client.get("/organization/recommendations").json()
    assert recs["basis"] and len(recs["startups"]) == 10
    scores = [s["score"] for s in recs["startups"]]
    assert scores == sorted(scores, reverse=True)
    [s01] = [s for s in recs["startups"] if s["id"] == "s01"]
    assert s01["on_platform"] and s01["applied"]
    assert s01["reason"].startswith("Aradığınız “")
    assert not any(s["on_platform"] for s in recs["startups"] if s["id"] != "s01")
